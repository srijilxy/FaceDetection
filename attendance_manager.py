from openpyxl import Workbook, load_workbook
import os
import datetime
from config import EXCEL_FILE

def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"
        ws.append(["Name", "Attendance", "Date", "Time"])
        wb.save(EXCEL_FILE)

def mark_attendance(name):
    init_excel()
    now = datetime.datetime.now()
    date_str = now.strftime("%d-%m-%y")
    time_str = now.strftime("%H:%M:%S")
    wb = load_workbook(EXCEL_FILE)
    ws = wb["Attendance"]
    names_today = [row[0].value for row in ws.iter_rows(min_row=2) if row.value == date_str]
    if name not in names_today:
        ws.append([name, "Present", date_str, time_str])
        wb.save(EXCEL_FILE)
